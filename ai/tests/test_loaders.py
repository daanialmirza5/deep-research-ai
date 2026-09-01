import io

import docx
import httpx
import pytest

from ai.loaders.base import LoaderError
from ai.loaders.csv_loader import CsvLoader
from ai.loaders.docx_loader import DocxLoader
from ai.loaders.markdown_loader import MarkdownLoader
from ai.loaders.pdf_loader import PdfLoader
from ai.loaders.registry import get_loader
from ai.loaders.text_loader import TextLoader
from ai.loaders.url_loader import UrlLoader
from ai.loaders.youtube_loader import YoutubeLoader, _extract_video_id


def _build_minimal_pdf(text: str) -> bytes:
    """A hand-built, minimally valid single-page PDF with one text run —
    pypdf has no page-writing API of its own (it's a reader/manipulator
    library), so this is the standard way to get real, parseable PDF bytes
    without pulling in a PDF-generation dependency just for tests."""
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/Resources<</Font<</F1 4 0 R>>>>"
        b"/MediaBox[0 0 200 200]/Contents 5 0 R>>",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    stream = f"BT /F1 24 Tf 10 100 Td ({text}) Tj ET".encode()
    objects.append(
        b"<</Length " + str(len(stream)).encode() + b">>\nstream\n" + stream + b"\nendstream"
    )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(str(index).encode() + b" 0 obj" + body + b"endobj\n")
    xref_offset = out.tell()
    count = len(objects) + 1
    out.write(b"xref\n")
    out.write(f"0 {count}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(b"trailer\n")
    out.write(f"<</Size {count}/Root 1 0 R>>\n".encode())
    out.write(b"startxref\n" + str(xref_offset).encode() + b"\n%%EOF")
    return out.getvalue()


@pytest.mark.asyncio
async def test_pdf_loader_extracts_text() -> None:
    text = await PdfLoader().load(_build_minimal_pdf("Hello World"))
    assert "Hello World" in text


@pytest.mark.asyncio
async def test_pdf_loader_rejects_a_url_source() -> None:
    with pytest.raises(LoaderError, match="expects raw file bytes"):
        await PdfLoader().load("not bytes")


@pytest.mark.asyncio
async def test_pdf_loader_rejects_garbage_bytes() -> None:
    with pytest.raises(LoaderError, match="could not read PDF"):
        await PdfLoader().load(b"this is not a pdf")


@pytest.mark.asyncio
async def test_docx_loader_extracts_paragraphs() -> None:
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = await DocxLoader().load(buffer.getvalue())

    assert "First paragraph." in text
    assert "Second paragraph." in text


@pytest.mark.asyncio
async def test_docx_loader_rejects_garbage_bytes() -> None:
    with pytest.raises(LoaderError, match="could not read DOCX"):
        await DocxLoader().load(b"this is not a docx")


@pytest.mark.asyncio
async def test_text_loader_decodes_utf8() -> None:
    text = await TextLoader().load(b"plain text content")
    assert text == "plain text content"


@pytest.mark.asyncio
async def test_text_loader_rejects_empty_file() -> None:
    with pytest.raises(LoaderError, match="empty"):
        await TextLoader().load(b"   ")


@pytest.mark.asyncio
async def test_markdown_loader_strips_formatting_syntax() -> None:
    text = await MarkdownLoader().load(b"# Heading\n\nSome **bold** text.")
    assert "Heading" in text
    assert "bold" in text
    assert "#" not in text
    assert "**" not in text


@pytest.mark.asyncio
async def test_csv_loader_formats_rows_as_key_value_pairs() -> None:
    csv_bytes = b"name,age\nAda,36\nGrace,85\n"
    text = await CsvLoader().load(csv_bytes)
    lines = text.splitlines()
    assert lines == ["name: Ada, age: 36", "name: Grace, age: 85"]


@pytest.mark.asyncio
async def test_csv_loader_rejects_a_header_only_csv() -> None:
    with pytest.raises(LoaderError, match="no rows"):
        await CsvLoader().load(b"name,age\n")


def test_youtube_video_id_extraction_handles_common_url_shapes() -> None:
    assert _extract_video_id("https://www.youtube.com/watch?v=dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert _extract_video_id("https://youtu.be/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert _extract_video_id("https://www.youtube.com/embed/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    with pytest.raises(LoaderError, match="could not find a video id"):
        _extract_video_id("https://example.com/not-a-video")


def _url_reachable(url: str) -> bool:
    try:
        return httpx.get(url, timeout=3.0, follow_redirects=True).status_code == 200
    except httpx.HTTPError:
        return False


@pytest.mark.skipif(
    not _url_reachable("https://example.com"), reason="example.com unreachable right now"
)
@pytest.mark.asyncio
async def test_url_loader_live_fetch_and_strip() -> None:
    text = await UrlLoader().load("https://example.com")
    assert "Example Domain" in text
    assert "<html" not in text.lower()


@pytest.mark.asyncio
async def test_youtube_loader_live_fetch_transcript() -> None:
    # A long-stable, widely-captioned public video — same "run live, skip on
    # unreachable" posture as ai/tools' tests (see test_tools.py's docstring),
    # since transcript availability/YouTube reachability isn't guaranteed
    # from every network.
    try:
        text = await YoutubeLoader().load("https://www.youtube.com/watch?v=jNQXAC9IdLA")
    except LoaderError as exc:
        pytest.skip(f"YouTube transcript unavailable right now: {exc}")
    assert len(text) > 0


def test_registry_returns_the_matching_loader_for_every_source_type() -> None:
    for source_type, loader_cls in (
        ("pdf", PdfLoader),
        ("docx", DocxLoader),
        ("txt", TextLoader),
        ("markdown", MarkdownLoader),
        ("csv", CsvLoader),
        ("url", UrlLoader),
        ("youtube", YoutubeLoader),
    ):
        assert isinstance(get_loader(source_type), loader_cls)


def test_registry_rejects_an_unknown_source_type() -> None:
    with pytest.raises(ValueError, match="no loader registered"):
        get_loader("carrier-pigeon")
